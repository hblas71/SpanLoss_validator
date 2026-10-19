import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

COLOR_NOKIA_BLUE = RGBColor(18, 101, 196)
COLOR_DARK_NAVY = RGBColor(15, 42, 74)
COLOR_RED = RGBColor(217, 30, 24)
COLOR_GREEN = RGBColor(39, 174, 96)
COLOR_GRAY = RGBColor(100, 100, 100)
HEX_LIGHT_RED = "FDEED9"
HEX_LIGHT_GRAY = "F4F6F8"
HEX_HEADER_BG = "0F2A4A"

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''<w:tcMar {nsdecls("w")}>
        <w:top w:w="{top}" w:type="dxa"/>
        <w:bottom w:w="{bottom}" w:type="dxa"/>
        <w:left w:w="{left}" w:type="dxa"/>
        <w:right w:w="{right}" w:type="dxa"/>
    </w:tcMar>''')
    tcPr.append(tcMar)

def build_docx_report(results, output_path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    excel_info = results.get("excel_info", {})
    nodo_a = excel_info.get("nodo_a", "Nodo A")
    nodo_b = excel_info.get("nodo_b", "Nodo B")

    p_title = doc.add_paragraph()
    r_title = p_title.add_run("INFORME TÉCNICO DE CARACTERIZACIÓN DE FIBRA ÓPTICA")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = COLOR_DARK_NAVY

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run(f"Vano: {nodo_a} ↔ {nodo_b} | Norma ITU-T G.652.D")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = COLOR_NOKIA_BLUE

    # 1. Resumen Ejecutivo
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Resumen Ejecutivo")
    r_h1.font.color.rgb = COLOR_DARK_NAVY

    headers = ["Hilo", "Longitud (km)", "Pérdida Vano (dB)", "Desbalance (dB)", "Pérdida Central (dB)", "Eventos NOK", "Dictamen"]
    t_exec = doc.add_table(rows=1, cols=len(headers))
    t_exec.alignment = WD_TABLE_ALIGNMENT.CENTER

    for c_idx, h_text in enumerate(headers):
        c = t_exec.cell(0, c_idx)
        set_cell_background(c, HEX_HEADER_BG)
        set_cell_margins(c, top=80, bottom=80, left=60, right=60)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    delta_loss = results.get("pair_symmetry", {}).get("delta_loss_db", 0.0)
    for t in results.get("traces_full_span", []):
        co = next((c for c in results.get("traces_central_exit", []) if c["direction"] == t["direction"]), {})
        row_cells = t_exec.add_row().cells
        vals = [
            t["direction"],
            f"{t['span_length_km']:.3f}",
            f"{t['total_loss_db']:.2f}",
            f"{delta_loss:.2f}",
            f"{co.get('loss_db', 0.0):.2f}",
            str(len(t.get("out_of_spec_events", []))),
            t["status"]
        ]
        for c_idx, val in enumerate(vals):
            c = row_cells[c_idx]
            bg = HEX_LIGHT_RED if t["status"] == "NO APTO" else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, top=60, bottom=60, left=60, right=60)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if val == "NO APTO":
                r.bold = True
                r.font.color.rgb = COLOR_RED

    # Dictamen Callout
    status = results.get("global_status", "NO APTO")
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(12)
    r_b = p_box.add_run(f"DICTAMEN GLOBAL: {status}\n")
    r_b.bold = True
    r_b.font.color.rgb = COLOR_GREEN if status == "APTO" else COLOR_RED
    for reas in results.get("verdict_reasons", []):
        r_r = p_box.add_run(f"• {reas}\n")
        r_r.font.size = Pt(9.5)

    # 2. Plan de Actuación
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. Plan de Actuación Priorizado")
    r_h2.font.color.rgb = COLOR_DARK_NAVY

    p_act = doc.add_paragraph()
    p_act.add_run(
        "P1 (Crítica) — Cabecera ODF: Limpieza microscópica y saneamiento de conectores en repartidor para reducir pérdida en salida de central < 2.00 dB.\n"
        "P2 (Alta) — Planta Externa: Refundir empalmes defectuosos por encima de 0.20 dB.\n"
        "P3 (Media) — Medición de Contraste: Repetir reflectometría bidireccional de comprobación."
    )

    p_disc = doc.add_paragraph()
    p_disc.paragraph_format.space_before = Pt(20)
    r_disc = p_disc.add_run("🤖 This is an AI generated response, it may not be fully accurate. Please verify the information independently.")
    r_disc.italic = True
    r_disc.font.size = Pt(8.5)
    r_disc.font.color.rgb = COLOR_GRAY

    doc.save(output_path)
