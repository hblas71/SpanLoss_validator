import streamlit as st
import os
import sys
import tempfile
import struct
import zipfile
import re
import json
import pandas as pd
import openpyxl
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# ==========================================
# 1. PARSER TELCORDIA .SOR (SR-4731 / GR-196)
# ==========================================
c_light = 299792458.0

def read_null_term_string(data, offset):
    end = data.find(b"\x00", offset)
    if end == -1:
        return "", len(data)
    try:
        val = data[offset:end].decode("utf-8", errors="replace").strip()
    except Exception:
        val = data[offset:end].decode("latin-1", errors="replace").strip()
    return val, end + 1

class SorParser:
    def __init__(self, filepath_or_bytes):
        if isinstance(filepath_or_bytes, (str, os.PathLike)) and os.path.exists(str(filepath_or_bytes)):
            with open(filepath_or_bytes, "rb") as f:
                self.data = f.read()
        elif isinstance(filepath_or_bytes, bytes):
            self.data = filepath_or_bytes
        else:
            raise ValueError("Entrada no valida para SorParser")

        self.metadata = {}
        self.params = {}
        self.events = []
        self.summary = {}
        self.map_size = 0

    def parse(self):
        data = self.data
        if len(data) < 12:
            raise ValueError("Archivo demasiado corto para formato SOR")

        tag = data[:4]
        if tag.startswith(b"Map"):
            map_ver, self.map_size, n_blocks = struct.unpack("<HIh", data[4:12])
        else:
            self.map_size = 256

        self._parse_gen_params()
        self._parse_fxd_params()
        self._parse_key_events()
        self._calculate_summary()

        return {
            "metadata": self.metadata,
            "params": self.params,
            "summary": self.summary,
            "events": self.events
        }

    def _parse_gen_params(self):
        data = self.data
        tag = b"GenParams\x00"
        pos = data.find(tag, max(12, self.map_size - 10))
        if pos == -1:
            pos = data.find(tag, 100)
        if pos == -1:
            return

        curr = pos + len(tag)
        cable_id, curr = read_null_term_string(data, curr)
        fiber_id, curr = read_null_term_string(data, curr)
        self.metadata["cable_id"] = cable_id
        self.metadata["fiber_id"] = fiber_id

        fiber_type, curr = read_null_term_string(data, curr)
        self.metadata["fiber_type"] = fiber_type

        if curr + 2 <= len(data):
            nom_wl = struct.unpack("<H", data[curr:curr+2])[0]
            self.params["nominal_wavelength_nm"] = nom_wl
            curr += 2

        loc_a, curr = read_null_term_string(data, curr)
        loc_b, curr = read_null_term_string(data, curr)
        self.metadata["originating_location"] = loc_a
        self.metadata["terminating_location"] = loc_b

        cable_code, curr = read_null_term_string(data, curr)
        self.metadata["cable_code"] = cable_code
        curr += 2
        if curr + 4 <= len(data):
            curr += 4

        operator, curr = read_null_term_string(data, curr)
        comments, curr = read_null_term_string(data, curr)
        self.metadata["operator"] = operator
        self.metadata["comments"] = comments

    def _parse_fxd_params(self):
        data = self.data
        tag = b"FxdParams\x00"
        pos = data.find(tag, max(12, self.map_size - 10))
        if pos == -1:
            pos = data.find(tag, 100)
        if pos == -1:
            return

        off = pos + len(tag)
        if off + 40 > len(data):
            return

        ts = struct.unpack("<I", data[off:off+4])[0]
        self.metadata["timestamp_unix"] = ts
        off += 4

        units = data[off:off+2].decode("ascii", errors="ignore").strip()
        self.params["units"] = units or "km"
        off += 2

        actual_wl, acq_off, acq_off_dist, n_pw, pw, spacing, n_pts, ior_val = struct.unpack("<HiiHHIIi", data[off:off+26])
        off += 26

        self.params["wavelength_nm"] = round(actual_wl / 10.0, 1)
        self.params["pulse_width_ns"] = pw
        self.params["sample_spacing_raw"] = spacing
        self.params["num_data_points"] = n_pts

        actual_ior = round(ior_val / 100000.0, 5) if ior_val > 100000 else 1.46500
        self.params["group_index_ior"] = actual_ior

        if off + 14 <= len(data):
            backscatter = struct.unpack("<h", data[off:off+2])[0]
            self.params["backscatter_coeff_db"] = round(-abs(backscatter / 10.0), 1)
            off += 2

            n_avg, avg_t, rng = struct.unpack("<HHI", data[off:off+8])
            off += 8
            rng_km = round((rng * 1e-10 * c_light) / (actual_ior * 1000.0), 2) if rng > 0 else 0.0
            self.params["range_km"] = rng_km

    def _parse_key_events(self):
        data = self.data
        tag = b"KeyEvents\x00"
        pos = data.find(tag, max(12, self.map_size - 10))
        if pos == -1:
            pos = data.find(tag, 100)
        if pos == -1:
            return

        off = pos + len(tag)
        if off + 2 > len(data):
            return

        num_events = struct.unpack("<H", data[off:off+2])[0]
        off += 2
        actual_ior = self.params.get("group_index_ior", 1.46500)

        for i in range(num_events):
            if off + 42 > len(data):
                break

            ev_num, prop_time, att_coeff, ev_loss, ev_refl = struct.unpack("<HIhhi", data[off:off+14])
            off += 14

            ev_code = data[off:off+6].decode("ascii", errors="replace")
            off += 6

            loss_tech = data[off:off+2].decode("ascii", errors="replace")
            off += 2

            m1, m2, m3, m4, m5 = struct.unpack("<iiiii", data[off:off+20])
            off += 20

            comment, off = read_null_term_string(data, off)

            dist_m = (prop_time * 1e-10 * c_light) / actual_ior
            dist_km = dist_m / 1000.0
            loss_db = round(ev_loss / 1000.0, 3)
            refl_db = round(ev_refl / 1000.0, 2) if ev_refl != 0 else None
            slope_db_km = round(att_coeff / 1000.0, 3) if att_coeff != 0 else None

            event_type = "Splice / Fusión"
            if ev_code.startswith("1") or (refl_db is not None and refl_db < 0):
                event_type = "Reflective / Conector"
            if "E" in ev_code or i == num_events - 1:
                event_type = "Fin de Fibra"

            self.events.append({
                "event_number": ev_num or (i + 1),
                "distance_m": round(dist_m, 2),
                "distance_km": round(dist_km, 4),
                "event_type": event_type,
                "event_code": ev_code,
                "loss_db": loss_db,
                "reflectance_db": refl_db,
                "slope_db_km": slope_db_km,
                "comment": comment
            })

    def _calculate_summary(self):
        if not self.events:
            return
        last_ev = self.events[-1]
        span_length_km = last_ev["distance_km"]

        total_loss_db = None
        data = self.data
        pos_jdsu = data.find(b"JDSU", max(100, self.map_size - 10))
        if pos_jdsu != -1:
            for i in range(len(data) - 5000, len(data) - 8):
                if data[i:i+4] == b"\x00\x00\xff?":
                    try:
                        v = struct.unpack("<d", data[i:i+8])[0]
                        if 1.5 <= v <= 35.0:
                            total_loss_db = round(v, 2)
                            break
                    except Exception:
                        pass

            if total_loss_db is None:
                for i in range(len(data) - 4000, len(data) - 8, 2):
                    try:
                        val = struct.unpack("<d", data[i:i+8])[0]
                        if 2.5 <= val <= 30.0 and val != 10.0:
                            total_loss_db = round(val, 2)
                            break
                    except Exception:
                        pass

        if total_loss_db is None:
            disc_loss = sum(ev["loss_db"] for ev in self.events if ev["loss_db"] > 0)
            fiber_loss = span_length_km * 0.22
            total_loss_db = round(disc_loss + fiber_loss, 2)

        avg_att = round(total_loss_db / span_length_km, 3) if span_length_km > 0 else 0.0

        bc = self.params.get("backscatter_coeff_db", -81.0)
        est_orl = round(abs(bc) - (total_loss_db * 1.1), 1)
        if est_orl < 25.0:
            est_orl = 28.5

        self.summary = {
            "span_length_km": round(span_length_km, 3),
            "span_length_m": round(last_ev["distance_m"], 1),
            "total_loss_db": round(total_loss_db, 2),
            "avg_attenuation_db_km": avg_att,
            "orl_db": est_orl,
            "event_count": len(self.events)
        }

# ==========================================
# 2. EXTRACTOR OLE DESDE EXCEL
# ==========================================
def extract_sor_from_excel(excel_path, output_dir, extract_dir):
    os.makedirs(output_dir, exist_ok=True)
    os.makedirs(extract_dir, exist_ok=True)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet_name = next((s for s in wb.sheetnames if s.strip().upper() == "INFORME"), None)

    nodo_a, nodo_b = "NodoA", "NodoB"
    if sheet_name:
        ws = wb[sheet_name]
        for r in range(1, 40):
            for c in range(1, 10):
                val = ws.cell(r, c).value
                if val and isinstance(val, str):
                    clean_val = val.strip().lower()
                    if clean_val == "nodo a":
                        cand = ws.cell(r, c+1).value or ws.cell(r+1, c).value
                        if cand:
                            nodo_a = str(cand).strip()
                    elif clean_val == "nodo b":
                        cand = ws.cell(r, c+1).value or ws.cell(r+1, c).value
                        if cand:
                            nodo_b = str(cand).strip()

    safe_nodo_a = re.sub(r'[/\\:*?"<>|]', '_', nodo_a)
    safe_nodo_b = re.sub(r'[/\\:*?"<>|]', '_', nodo_b)
    zip_name = f"OTDR_{safe_nodo_a}-{safe_nodo_b}.zip"
    zip_path = os.path.join(output_dir, zip_name)

    extracted = {}
    with zipfile.ZipFile(excel_path, "r") as z:
        for name in sorted(z.namelist()):
            if name.startswith("xl/embeddings/oleObject") and name.endswith(".bin"):
                data = z.read(name)
                first_sor = data.find(b".sor")
                if first_sor == -1:
                    continue
                start_pos = data.rfind(b"\xff\xff\xff\xff", 0, first_sor) + 4
                if start_pos < 4:
                    continue

                pos = start_pos
                total_size = struct.unpack("<I", data[pos:pos+4])[0]
                pos += 4
                flags = struct.unpack("<H", data[pos:pos+2])[0]
                pos += 2

                fn_end = data.find(b"\x00", pos)
                filename = data[pos:fn_end].decode("latin1")
                pos = fn_end + 1

                sp_end = data.find(b"\x00", pos)
                src_path = data[pos:sp_end].decode("latin1")
                pos = sp_end + 1
                pos += 4

                tp_len = struct.unpack("<I", data[pos:pos+4])[0]
                if data[pos+4:pos+4+tp_len].endswith(b"\x00"):
                    pos = pos + 4 + tp_len
                else:
                    tp_end = data.find(b"\x00", pos)
                    pos = tp_end + 1

                data_size = struct.unpack("<I", data[pos:pos+4])[0]
                pos += 4
                file_bytes = data[pos:pos+data_size]

                if file_bytes.startswith(b"Map\x00") or b"Cksum\x00" in file_bytes[-20:]:
                    extracted[filename] = file_bytes

    if not extracted:
        raise ValueError(f"No se encontraron trazas .sor validas dentro de {excel_path}")

    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for fn, b in extracted.items():
            zout.writestr(fn, b)

    saved_files = []
    for fn, b in extracted.items():
        dst = os.path.join(extract_dir, fn)
        with open(dst, "wb") as f:
            f.write(b)
        saved_files.append(dst)

    return {
        "nodo_a": nodo_a,
        "nodo_b": nodo_b,
        "zip_path": zip_path,
        "saved_file_paths": saved_files
    }

# ==========================================
# 3. MOTOR DE AUDITORIA G.652.D
# ==========================================
THRESHOLDS = {
    "splice_loss_max_db": 0.20,
    "connector_loss_max_db": 0.50,
    "central_office_loss_max_db": 2.00,
    "pair_delta_length_pct_max": 0.05,
    "pair_delta_loss_db_max": 0.50
}

def parse_filename_meta(filename):
    base = os.path.basename(filename)
    name_no_ext = os.path.splitext(base)[0]
    tokens = re.split(r'[-_ ]+', name_no_ext)
    direction = "UNKNOWN"
    for t in tokens:
        up = t.upper()
        if up in ["LI", "RX"]:
            direction = "LI"
            break
        elif up in ["LO", "TX"]:
            direction = "LO"
            break
    is_10km = any(k in name_no_ext.lower() for k in ["10k", "10km"])
    return {"filename": base, "direction": direction, "is_10km": is_10km}

def run_analysis(sor_files, excel_info=None):
    traces = []
    for fp in sor_files:
        meta = parse_filename_meta(fp)
        p = SorParser(fp)
        data = p.parse()
        data["meta"] = meta
        traces.append(data)

    traces_10k = [t for t in traces if t["meta"]["is_10km"]]
    traces_full = [t for t in traces if not t["meta"]["is_10km"]]
    if not traces_full and traces_10k:
        traces_full = traces_10k

    res = {
        "excel_info": excel_info,
        "traces_full_span": [],
        "traces_central_exit": [],
        "pair_symmetry": {},
        "global_status": "APTO",
        "verdict_reasons": []
    }

    li_full = next((t for t in traces_full if t["meta"]["direction"] == "LI"), None)
    lo_full = next((t for t in traces_full if t["meta"]["direction"] == "LO"), None)

    for t in traces_full:
        summary = t.get("summary", {})
        meta = t["meta"]
        events = t.get("events", [])
        out_events = []
        for ev in events:
            ev_t = ev.get("event_type", "")
            l = ev.get("loss_db", 0.0)
            if "Splice" in ev_t and l > THRESHOLDS["splice_loss_max_db"]:
                out_events.append({"event": ev, "reason": f"Empalme ({l} dB > {THRESHOLDS['splice_loss_max_db']} dB)"})
            elif "Reflective" in ev_t and l > THRESHOLDS["connector_loss_max_db"]:
                out_events.append({"event": ev, "reason": f"Conector ({l} dB > {THRESHOLDS['connector_loss_max_db']} dB)"})

        st_thr = "APTO" if len(out_events) == 0 else "NO APTO"
        if st_thr == "NO APTO":
            res["global_status"] = "NO APTO"
            res["verdict_reasons"].append(f"Hilo {meta['direction']}: {len(out_events)} eventos fuera de norma")

        res["traces_full_span"].append({
            "filename": meta["filename"],
            "direction": meta["direction"],
            "span_length_km": summary.get("span_length_km", 0.0),
            "total_loss_db": summary.get("total_loss_db", 0.0),
            "out_of_spec_events": out_events,
            "status": st_thr
        })

    if li_full and lo_full:
        l_li = li_full.get("summary", {}).get("span_length_km", 0.0)
        l_lo = lo_full.get("summary", {}).get("span_length_km", 0.0)
        loss_li = li_full.get("summary", {}).get("total_loss_db", 0.0)
        loss_lo = lo_full.get("summary", {}).get("total_loss_db", 0.0)

        delta_len = abs(l_li - l_lo)
        delta_loss = abs(loss_li - loss_lo)
        max_l = max(l_li, l_lo)
        delta_pct = (delta_len / max_l * 100.0) if max_l > 0 else 0.0

        if delta_pct > (THRESHOLDS["pair_delta_length_pct_max"] * 100):
            res["global_status"] = "NO APTO"
            res["verdict_reasons"].append(f"Desbalance longitud {delta_pct:.2f}% > 5%")
        if delta_loss > THRESHOLDS["pair_delta_loss_db_max"]:
            res["global_status"] = "NO APTO"
            res["verdict_reasons"].append(f"Desbalance de atenuación {delta_loss:.2f} dB > {THRESHOLDS['pair_delta_loss_db_max']} dB")

        res["pair_symmetry"] = {
            "delta_length_km": delta_len,
            "delta_loss_db": delta_loss
        }

    for t in traces_10k:
        events = t.get("events", [])
        meta = t["meta"]
        co_events = [e for e in events if 25.0 <= e.get("distance_m", 0.0) <= 200.0]
        co_loss = sum(e.get("loss_db", 0.0) for e in co_events if e.get("loss_db", 0.0) > 0)
        st_co = "APTO" if co_loss <= THRESHOLDS["central_office_loss_max_db"] else "NO APTO"
        if st_co == "NO APTO":
            res["global_status"] = "NO APTO"
            res["verdict_reasons"].append(f"Salida Central {meta['direction']}: Pérdida de {co_loss:.2f} dB > 2.00 dB")

        res["traces_central_exit"].append({
            "direction": meta["direction"],
            "loss_db": co_loss,
            "status": st_co
        })

    return res

# ==========================================
# 4. GENERADOR WORD
# ==========================================
def build_docx(results, output_path):
    doc = Document()
    excel_info = results.get("excel_info") or {}
    nodo_a = excel_info.get("nodo_a", "Nodo A")
    nodo_b = excel_info.get("nodo_b", "Nodo B")

    p = doc.add_paragraph()
    r = p.add_run("INFORME DE CARACTERIZACIÓN OTDR (G.652.D)")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(15, 42, 74)

    p2 = doc.add_paragraph(f"Vano: {nodo_a} ↔ {nodo_b}")
    p2.runs[0].font.color.rgb = RGBColor(18, 101, 196)

    st_glob = results.get("global_status", "NO APTO")
    p_st = doc.add_paragraph()
    r_st = p_st.add_run(f"DICTAMEN GLOBAL: {st_glob}\n")
    r_st.bold = True
    r_st.font.color.rgb = RGBColor(39, 174, 96) if st_glob == "APTO" else RGBColor(217, 30, 24)
    for reason in results.get("verdict_reasons", []):
        p_st.add_run(f"• {reason}\n")

    doc.add_heading("Resumen de Par de Fibras", level=2)
    tbl = doc.add_table(rows=1, cols=6)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdrs = ["Hilo", "Longitud (km)", "Pérdida (dB)", "Desbalance", "Salida Central", "Dictamen"]
    for i, h in enumerate(hdrs):
        tbl.cell(0, i).paragraphs[0].add_run(h).bold = True

    delta_loss = results.get("pair_symmetry", {}).get("delta_loss_db", 0.0)
    for t in results.get("traces_full_span", []):
        co = next((c for c in results.get("traces_central_exit", []) if c["direction"] == t["direction"]), {})
        row = tbl.add_row().cells
        row[0].paragraphs[0].add_run(t["direction"])
        row[1].paragraphs[0].add_run(f"{t['span_length_km']:.3f}")
        row[2].paragraphs[0].add_run(f"{t['total_loss_db']:.2f}")
        row[3].paragraphs[0].add_run(f"{delta_loss:.2f} dB")
        row[4].paragraphs[0].add_run(f"{co.get('loss_db', 0.0):.2f} dB")
        r_d = row[5].paragraphs[0].add_run(t["status"])
        if t["status"] == "NO APTO":
            r_d.bold = True
            r_d.font.color.rgb = RGBColor(217, 30, 24)

    doc.save(output_path)

# ==========================================
# 5. INTERFAZ STREAMLIT
# ==========================================
st.set_page_config(page_title="Nokia OTDR Analyzer", page_icon="📶", layout="wide")
st.title("📶 Caracterización Técnica de Fibras Ópticas (OTDR)")
st.caption("Norma ITU-T G.652.D | Recorrido Completo y Salida de Central")

up = st.file_uploader("Arrastra la plantilla Excel (.xlsm/.xlsx), ZIP o fichero .sor:", type=["xlsm", "xlsx", "zip", "sor"])
if up:
    with tempfile.TemporaryDirectory() as td:
        inp = os.path.join(td, up.name)
        with open(inp, "wb") as f:
            f.write(up.getbuffer())

        out_d = os.path.join(td, "out")
        ext_d = os.path.join(td, "ext")
        os.makedirs(out_d, exist_ok=True)
        os.makedirs(ext_d, exist_ok=True)

        excel_info = None
        sor_files = []

        if inp.lower().endswith(('.xlsm', '.xlsx')):
            excel_info = extract_sor_from_excel(inp, out_d, ext_d)
            sor_files = excel_info["saved_file_paths"]
        elif inp.lower().endswith(".zip"):
            with zipfile.ZipFile(inp, "r") as z:
                z.extractall(ext_d)
                for root, _, fs in os.walk(ext_d):
                    for fn in fs:
                        if fn.lower().endswith(".sor"):
                            sor_files.append(os.path.join(root, fn))
        elif inp.lower().endswith(".sor"):
            sor_files = [inp]

        res = run_analysis(sor_files, excel_info=excel_info)
        nodo_a = (excel_info or {}).get("nodo_a", "NodoA")
        nodo_b = (excel_info or {}).get("nodo_b", "NodoB")

        st.subheader(f"Vano: {nodo_a} ↔ {nodo_b}")
        if res["global_status"] == "APTO":
            st.success("### 🟢 DICTAMEN GLOBAL: VANO APTO")
        else:
            st.error("### 🔴 DICTAMEN GLOBAL: VANO NO APTO (NO OK)")
            for r in res.get("verdict_reasons", []):
                st.write(f"- ❌ {r}")

        t1, t2 = st.tabs(["📊 Resumen", "📥 Descargas"])
        with t1:
            rows = []
            for t in res.get("traces_full_span", []):
                co = next((c for c in res.get("traces_central_exit", []) if c["direction"] == t["direction"]), {})
                rows.append({
                    "Hilo": t["direction"],
                    "Longitud (km)": f"{t['span_length_km']:.3f}",
                    "Pérdida Vano (dB)": f"{t['total_loss_db']:.2f}",
                    "Desbalance Par": f"{res.get('pair_symmetry', {}).get('delta_loss_db', 0):.2f} dB",
                    "Pérdida Central": f"{co.get('loss_db', 0):.2f} dB",
                    "Eventos NOK": len(t.get("out_of_spec_events", [])),
                    "Dictamen": t["status"]
                })
            st.dataframe(pd.DataFrame(rows), use_container_width=True)

        with t2:
            docx_name = f"Informe_OTDR_{nodo_a}-{nodo_b}.docx"
            docx_p = os.path.join(out_d, docx_name)
            build_docx(res, docx_p)
            with open(docx_p, "rb") as f_d:
                st.download_button("📄 Descargar Informe Word (.docx)", f_d, file_name=docx_name)

            if excel_info and os.path.exists(excel_info.get("zip_path", "")):
                with open(excel_info["zip_path"], "rb") as f_z:
                    st.download_button("📦 Descargar Trazas .sor (ZIP)", f_z, file_name=os.path.basename(excel_info["zip_path"]))
